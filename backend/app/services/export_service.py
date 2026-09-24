from __future__ import annotations

import json
from pathlib import Path

from docx import Document


def export_txt(text: str, output_path: Path) -> Path:
    output_path.write_text(text or "", encoding="utf-8")
    return output_path


def export_json(payload: dict, output_path: Path) -> Path:
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return output_path


def export_docx(text: str, output_path: Path, title: str = "Meeting Transcript") -> Path:
    document = Document()
    document.add_heading(title, level=1)
    for paragraph in (text or "").split("\n"):
        document.add_paragraph(paragraph)
    document.save(output_path)
    return output_path
